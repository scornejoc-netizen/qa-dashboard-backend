"""
Importa el Excel CRONOGRAMA JIRA al admin de Django.

Uso:
    python manage.py import_excel_cronograma "C:/Users/scornejoc/Downloads/CRONOGRAMA JIRA.xlsx"
    python manage.py import_excel_cronograma "ruta.xlsx" --dry-run

Estructura esperada del Excel (a partir de la fila 3, con headers en fila 2):
    A: ESPACIO          - módulo funcional (string)
    B: FECHA INICIO     - datetime
    C: FECHA FINAL      - datetime
    D: ASIGNADO         - nombre del dev (o "NADIE" / "TODOS")
    E: LINK             - URL a Jira
    F: ESTADO           - "POR HACER" / "TERMINADO"
    G: DASHBOARD        - "TRACKED" / "UNTRACKED" (ignorado en la importación)
    H: DOCUMENTADO      - "SÍ" / "NO"

Comportamiento:
- Filas con ASIGNADO en {"NADIE", "TODOS"} se saltan.
- Filas con ESPACIO vacío se saltan.
- Idempotente: matchea por (space, developer.full_name) y actualiza si ya existe.
- Crea desarrolladores que no existan en la DB (con nivel 'mid' por defecto — Stephano
  los reasigna después).
"""
from datetime import date

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

try:
    import openpyxl
except ImportError as e:
    raise CommandError(
        "openpyxl no está instalado. Corre: pip install openpyxl"
    ) from e

from developers.models import Developer, Requirement


STATUS_MAP = {
    'POR HACER': 'todo',
    'TERMINADO': 'done',
    'EN DESARROLLO': 'in_progress',
    'EN QA': 'qa',
}

SKIP_ASSIGNEES = {'NADIE', 'TODOS', '', None}


def _normalize_date(value):
    """Convierte un datetime de openpyxl en date, o devuelve None."""
    if value is None:
        return None
    if hasattr(value, 'date'):
        return value.date()
    return value


def _truthy_yes(value):
    """'SÍ', 'SI', 'YES' → True; cualquier otra cosa → False."""
    if value is None:
        return False
    s = str(value).strip().upper().replace('Í', 'I')
    return s in {'SI', 'YES', 'TRUE', '1', 'OK'}


class Command(BaseCommand):
    help = 'Importa requerimientos desde el Excel CRONOGRAMA JIRA al admin de Django.'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str, help='Ruta absoluta al archivo .xlsx')
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Lee y reporta qué pasaría, pero no guarda nada en la DB.',
        )
        parser.add_argument(
            '--sheet',
            type=str,
            default='Hoja 1',
            help='Nombre de la hoja a leer (default: "Hoja 1").',
        )
        parser.add_argument(
            '--start-row',
            type=int,
            default=3,
            help='Fila desde donde empiezan los datos (default: 3 — fila 1 es título, fila 2 headers).',
        )

    def handle(self, *args, **options):
        path = options['excel_path']
        dry_run = options['dry_run']
        sheet_name = options['sheet']
        start_row = options['start_row']

        self.stdout.write(self.style.NOTICE(f"\nAbriendo Excel: {path}"))
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Archivo no encontrado: {path}")

        if sheet_name not in wb.sheetnames:
            raise CommandError(
                f"La hoja '{sheet_name}' no existe. Hojas disponibles: {wb.sheetnames}"
            )

        sh = wb[sheet_name]
        rows = list(sh.iter_rows(min_row=start_row, max_row=sh.max_row, values_only=True))
        self.stdout.write(f"Filas leídas: {len(rows)}\n")

        stats = {
            'skipped_empty': 0,
            'skipped_invalid_assignee': 0,
            'devs_created': 0,
            'devs_existing': 0,
            'reqs_created': 0,
            'reqs_updated': 0,
        }

        # Procesamiento (envuelto en transaction si no es dry-run)
        ctx = transaction.atomic() if not dry_run else _NoopContext()
        with ctx:
            for idx, row in enumerate(rows, start=start_row):
                self._process_row(row, idx, stats, dry_run)

            if dry_run:
                self.stdout.write(self.style.WARNING("\n[DRY-RUN] Nada se guardó en la DB.\n"))

        self._print_summary(stats, dry_run)

    def _process_row(self, row, row_idx, stats, dry_run):
        # Mapeo de columnas (índices 0-based)
        space = row[0]
        started_at = _normalize_date(row[1])
        delivered_at = _normalize_date(row[2])
        assignee = row[3]
        jira_url = row[4]
        estado = row[5]
        # row[6] es DASHBOARD — ignorado
        documentado = row[7] if len(row) > 7 else None

        # Skip filas vacías
        if space is None and assignee is None:
            stats['skipped_empty'] += 1
            return

        if space is None:
            self.stdout.write(self.style.WARNING(
                f"  Fila {row_idx}: ESPACIO vacío, se salta."
            ))
            stats['skipped_empty'] += 1
            return

        # Skip "NADIE" / "TODOS"
        assignee_str = str(assignee).strip() if assignee is not None else ''
        if assignee_str.upper() in {'NADIE', 'TODOS'} or not assignee_str:
            self.stdout.write(self.style.WARNING(
                f"  Fila {row_idx}: asignado='{assignee_str}' inválido, se salta."
            ))
            stats['skipped_invalid_assignee'] += 1
            return

        # Crear/obtener Developer
        if dry_run:
            try:
                dev = Developer.objects.get(full_name=assignee_str)
                stats['devs_existing'] += 1
            except Developer.DoesNotExist:
                self.stdout.write(f"  [DRY] Crearía dev: {assignee_str}")
                stats['devs_created'] += 1
                dev = None
        else:
            dev, dev_created = Developer.objects.get_or_create(
                full_name=assignee_str,
                defaults={'level': 'mid', 'active': True},
            )
            if dev_created:
                stats['devs_created'] += 1
                self.stdout.write(self.style.SUCCESS(f"  Dev creado: {assignee_str} (nivel mid por defecto)"))
            else:
                stats['devs_existing'] += 1

        # Mapeo de estado
        estado_upper = str(estado).strip().upper() if estado else 'POR HACER'
        status = STATUS_MAP.get(estado_upper, 'todo')

        # Para "POR HACER" la fecha final del Excel no debe poblarse como delivered_at
        if status != 'done':
            delivered_at = None

        # created_at: el Excel no lo tiene, usamos started_at o hoy
        created_at = started_at or date.today()

        # Datos del requerimiento
        title = str(space).strip()
        space_clean = str(space).strip()
        is_documented = _truthy_yes(documentado)

        if dry_run:
            self.stdout.write(
                f"  [DRY] Req: space={space_clean}, dev={assignee_str}, "
                f"status={status}, started={started_at}, delivered={delivered_at}, "
                f"documented={is_documented}"
            )
            return

        # Idempotencia: matchear por (space + dev incluido en el M2M)
        existing = Requirement.objects.filter(space=space_clean, developers=dev).first()
        fields = {
            'title': title,
            'status': status,
            'type': 'feature',  # tipo no está en el Excel
            'created_at': created_at,
            'started_at': started_at,
            'delivered_at': delivered_at,
            'jira_url': jira_url or '',
            'is_documented': is_documented,
        }
        if existing:
            for k, v in fields.items():
                setattr(existing, k, v)
            existing.save()
            req = existing
            created = False
        else:
            req = Requirement.objects.create(space=space_clean, **fields)
            req.developers.add(dev)
            created = True

        if created:
            stats['reqs_created'] += 1
            self.stdout.write(self.style.SUCCESS(
                f"  [+] Req creado: {req.code} | {space_clean} -> {dev.full_name}"
            ))
        else:
            stats['reqs_updated'] += 1
            self.stdout.write(
                f"  [~] Req actualizado: {req.code} | {space_clean} -> {dev.full_name}"
            )

    def _print_summary(self, stats, dry_run):
        prefix = "[DRY-RUN] " if dry_run else ""
        self.stdout.write(self.style.NOTICE(f"\n{prefix}--- Resumen ---"))
        self.stdout.write(f"  Filas saltadas (vacias):                {stats['skipped_empty']}")
        self.stdout.write(f"  Filas saltadas (NADIE/TODOS):           {stats['skipped_invalid_assignee']}")
        self.stdout.write(f"  Desarrolladores creados:                {stats['devs_created']}")
        self.stdout.write(f"  Desarrolladores ya existentes:          {stats['devs_existing']}")
        self.stdout.write(self.style.SUCCESS(f"  Requerimientos creados:                 {stats['reqs_created']}"))
        self.stdout.write(f"  Requerimientos actualizados:            {stats['reqs_updated']}")
        self.stdout.write(self.style.NOTICE("----------------\n"))

        if not dry_run and stats['devs_created'] > 0:
            self.stdout.write(self.style.WARNING(
                "Los desarrolladores nuevos se crearon con nivel 'mid' por defecto.\n"
                "Entra al admin y reasigna el nivel correcto (junior/senior) cuando puedas.\n"
            ))


class _NoopContext:
    """Context manager que no hace nada — para dry-run."""
    def __enter__(self): return self
    def __exit__(self, *args): return False
